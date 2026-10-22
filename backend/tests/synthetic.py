"""Synthetic test documents generated in-memory. No real customer data."""

from __future__ import annotations

import io
import struct
import zipfile
import zlib


def pdf_with_text(lines: list[str]) -> bytes:
    content = (
        "BT /F1 12 Tf 72 720 Td 14 TL "
        + " ".join(f"({ln.replace('(', '[').replace(')', ']')}) Tj T*" for ln in lines)
        + " ET"
    )
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def blank_pdf() -> bytes:
    return pdf_with_text([])


def png(width: int = 2, height: int = 2) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def jpeg() -> bytes:
    # Minimal JFIF header + EOI is enough for magic-byte detection tests.
    return bytes.fromhex("ffd8ffe000104a46494600010100000100010000ffd9")


def docx(paragraphs: list[str]) -> bytes:
    import docx as _docx

    d = _docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx(rows: list[list[object]], sheets: int = 1) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for i in range(1, sheets):
        wb.create_sheet(f"S{i}").append([f"sheet {i}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx(texts: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for t in texts:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def with_zip_member(data: bytes, name: str, content: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, content)
    return buf.getvalue()


def replace_zip_member(data: bytes, name: str, content: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, content if item.filename == name else src.read(item.filename))
    return buf.getvalue()


def zip_bomb_like() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", b"<Types/>")
        z.writestr("word/document.xml", b"\x00" * (60 * 1024 * 1024))
    return buf.getvalue()
