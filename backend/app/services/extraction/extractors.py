"""Format-specific extractors. No macros or embedded code are ever executed.

XML safety: python-docx and python-pptx parse with `resolve_entities=False`; openpyxl uses
defusedxml when installed (it is a declared dependency). Tests assert XXE payloads are inert.
"""

from __future__ import annotations

import csv
import io
import logging
import os
import shutil
import subprocess
import tempfile

from app.services.extraction.base import (
    DocumentExtractor,
    ExtractionLimits,
    ExtractionResult,
    TextAccumulator,
)
from app.services.file_validation import decode_text

log = logging.getLogger(__name__)
MIN_TEXT_CHARS = 20


def _finish(
    acc: TextAccumulator, fmt: str, count: int | None, warnings: list[str]
) -> ExtractionResult:
    if acc.truncated:
        warnings.append("text truncated at character limit")
    text = acc.text
    if len(text) < MIN_TEXT_CHARS:
        return ExtractionResult("empty", text, fmt, count, [*warnings, "little or no text found"])
    status = "partial" if warnings else "success"
    return ExtractionResult(status, text, fmt, count, warnings)  # type: ignore[arg-type]


class PdfExtractor(DocumentExtractor):
    kinds = ("pdf",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        from pypdf import PdfReader

        warnings: list[str] = []
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    return ExtractionResult(
                        "failed", format="pdf", warnings=["PDF is password protected"]
                    )
            except Exception:
                return ExtractionResult(
                    "failed", format="pdf", warnings=["PDF is password protected"]
                )
        pages = len(reader.pages)
        if pages > limits.max_pages:
            warnings.append(f"only the first {limits.max_pages} of {pages} pages were read")
        acc = TextAccumulator(limits.max_chars)
        for i, page in enumerate(reader.pages):
            if i >= limits.max_pages:
                break
            try:
                if not acc.add(page.extract_text() or ""):
                    break
            except Exception as exc:
                warnings.append(f"page {i + 1} unreadable ({type(exc).__name__})")
        result = _finish(acc, "pdf", pages, warnings)
        if result.status == "empty":
            result.warnings.append("PDF has no text layer (scanned?)")
        return result


class DocxExtractor(DocumentExtractor):
    kinds = ("docx",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        import docx

        document = docx.Document(io.BytesIO(data))
        acc = TextAccumulator(limits.max_chars)
        for p in document.paragraphs:
            if not acc.add(p.text):
                break
        for table in document.tables:
            for row in table.rows:
                if not acc.add(" | ".join(c.text.strip() for c in row.cells)):
                    break
        return _finish(acc, "docx", None, [])


class XlsxExtractor(DocumentExtractor):
    kinds = ("xlsx",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        import openpyxl

        warnings: list[str] = []
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = wb.sheetnames
            if len(sheets) > limits.max_sheets:
                warnings.append(f"only the first {limits.max_sheets} of {len(sheets)} sheets read")
            acc = TextAccumulator(limits.max_chars)
            for name in sheets[: limits.max_sheets]:
                acc.add(f"# Sheet: {name}")
                for r, row in enumerate(wb[name].iter_rows(values_only=True)):
                    if r >= limits.max_rows_per_sheet:
                        warnings.append(
                            f"sheet {name!r} truncated at {limits.max_rows_per_sheet} rows"
                        )
                        break
                    cells = [str(v) for v in row if v is not None]
                    if cells and not acc.add(" | ".join(cells)):
                        break
                if acc.truncated:
                    break
            return _finish(acc, "xlsx", len(sheets), warnings)
        finally:
            wb.close()


class XlsExtractor(DocumentExtractor):
    kinds = ("xls",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        import xlrd

        warnings: list[str] = []
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            n = book.nsheets
            if n > limits.max_sheets:
                warnings.append(f"only the first {limits.max_sheets} of {n} sheets read")
            acc = TextAccumulator(limits.max_chars)
            for idx in range(min(n, limits.max_sheets)):
                sheet = book.sheet_by_index(idx)
                acc.add(f"# Sheet: {sheet.name}")
                for r in range(min(sheet.nrows, limits.max_rows_per_sheet)):
                    cells = [str(v) for v in sheet.row_values(r) if v not in ("", None)]
                    if cells and not acc.add(" | ".join(cells)):
                        break
                book.unload_sheet(idx)
            return _finish(acc, "xls", n, warnings)
        finally:
            book.release_resources()


class PptxExtractor(DocumentExtractor):
    kinds = ("pptx",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        from pptx import Presentation

        warnings: list[str] = []
        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)
        if len(slides) > limits.max_slides:
            warnings.append(f"only the first {limits.max_slides} of {len(slides)} slides read")
        acc = TextAccumulator(limits.max_chars)
        for slide in slides[: limits.max_slides]:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame:
                    acc.add(shape.text_frame.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        acc.add(" | ".join(c.text for c in row.cells))
            if acc.truncated:
                break
        return _finish(acc, "pptx", len(slides), warnings)


class TextExtractor(DocumentExtractor):
    kinds = ("txt",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        text = decode_text(data)
        if text is None:
            return ExtractionResult("failed", format="txt", warnings=["not valid text"])
        acc = TextAccumulator(limits.max_chars)
        acc.add(text)
        return _finish(acc, "txt", None, [])


class CsvExtractor(DocumentExtractor):
    kinds = ("csv",)

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        text = decode_text(data)
        if text is None:
            return ExtractionResult("failed", format="csv", warnings=["not valid text"])
        acc = TextAccumulator(limits.max_chars)
        rows = 0
        warnings: list[str] = []
        try:
            for row in csv.reader(io.StringIO(text)):
                rows += 1
                if rows > limits.max_rows_per_sheet:
                    warnings.append(f"truncated at {limits.max_rows_per_sheet} rows")
                    break
                if not acc.add(" | ".join(row)):
                    break
        except csv.Error as exc:
            warnings.append(f"CSV parse problem: {exc}")
        return _finish(acc, "csv", None, warnings)


class ImageExtractor(DocumentExtractor):
    """No local OCR (keeps the image small). Images rely on AI vision when enabled."""

    kinds = ("jpeg", "png", "webp")

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        return ExtractionResult(
            "unsupported",
            format="image",
            warnings=["no local OCR; text read via AI vision if enabled"],
        )


class LegacyOfficeExtractor(DocumentExtractor):
    """.doc/.ppt: converted with headless LibreOffice only when LEGACY_CONVERTER=soffice."""

    kinds = ("doc", "ppt")

    def __init__(self, converter: str = "none") -> None:
        self.converter = converter

    def extract(self, data: bytes, limits: ExtractionLimits) -> ExtractionResult:
        if self.converter != "soffice" or not shutil.which("soffice"):
            return ExtractionResult(
                "unsupported",
                format="legacy-office",
                warnings=["legacy .doc/.ppt text extraction not enabled; original stored"],
            )
        with tempfile.TemporaryDirectory(prefix="fc-legacy-") as tmp:
            src = os.path.join(tmp, "input.bin")
            with open(src, "wb") as fh:
                fh.write(data)
            subprocess.run(  # noqa: S603 - fixed argv, no shell
                [
                    "soffice",
                    "--headless",
                    "--norestore",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    tmp,
                    src,
                ],
                check=False,
                capture_output=True,
                timeout=60,
                env={**os.environ, "HOME": tmp},
            )
            out = os.path.join(tmp, "input.txt")
            if not os.path.exists(out):
                return ExtractionResult(
                    "failed", format="legacy-office", warnings=["conversion failed"]
                )
            with open(out, "rb") as fh:
                text = decode_text(fh.read()) or ""
        acc = TextAccumulator(limits.max_chars)
        acc.add(text)
        return _finish(acc, "legacy-office", None, ["converted with LibreOffice"])


def extractor_for(kind: str, legacy_converter: str = "none") -> DocumentExtractor | None:
    table: dict[str, DocumentExtractor] = {
        "pdf": PdfExtractor(),
        "docx": DocxExtractor(),
        "xlsx": XlsxExtractor(),
        "xls": XlsExtractor(),
        "pptx": PptxExtractor(),
        "txt": TextExtractor(),
        "csv": CsvExtractor(),
        "jpeg": ImageExtractor(),
        "png": ImageExtractor(),
        "webp": ImageExtractor(),
        "doc": LegacyOfficeExtractor(legacy_converter),
        "ppt": LegacyOfficeExtractor(legacy_converter),
    }
    return table.get(kind)
